"""웹 화면 통합 테스트 — 로그인, 매트릭스, 팝업, 검색, 업로드→검토→반영."""

from __future__ import annotations

import re

import pytest

from tests.conftest import FIXTURES, requires_db

pytestmark = requires_db


@pytest.fixture()
def client(db_session):
    from fastapi.testclient import TestClient

    from app.main import app

    with TestClient(app, follow_redirects=False) as test_client:
        yield test_client


@pytest.fixture()
def admin(db_session):
    from app.models import AppUser
    from app.services.auth import hash_password

    user = AppUser(
        username="tester",
        display_name="테스터",
        password_hash=hash_password("verysecret123"),
        role="admin",
    )
    db_session.add(user)
    db_session.commit()
    return user


@pytest.fixture()
def auth(client, admin):
    response = client.post(
        "/login", data={"username": "tester", "password": "verysecret123", "next": "/"}
    )
    assert response.status_code == 303
    return client


def upload(auth_client, name: str):
    with (FIXTURES / name).open("rb") as handle:
        return auth_client.post("/upload", files={"file": (name, handle)})


class TestAuthGate:
    def test_root_redirects_to_login(self, client):
        response = client.get("/")
        assert response.status_code == 303
        assert "/login" in response.headers["location"]

    def test_wrong_password_is_rejected(self, client, admin):
        response = client.post(
            "/login", data={"username": "tester", "password": "nope", "next": "/"}
        )
        assert response.status_code == 401
        assert "올바르지 않습니다" in response.text

    def test_login_sets_cookie(self, auth):
        assert auth.cookies.get("contacts_session")

    def test_robots_blocks_indexing(self, client):
        assert "Disallow: /" in client.get("/robots.txt").text

    def test_security_headers(self, client):
        response = client.get("/login")
        assert response.headers["X-Robots-Tag"].startswith("noindex")
        assert response.headers["X-Frame-Options"] == "DENY"
        assert response.headers["Cache-Control"] == "no-store"


class TestEmptyState:
    def test_index_without_data(self, auth):
        response = auth.get("/")
        assert response.status_code == 200
        assert "아직 등록된 자료가 없습니다" in response.text


class TestFullCycle:
    @pytest.fixture()
    def loaded(self, auth):
        response = upload(auth, "sample_desk_matrix.xlsx")
        assert response.status_code == 303, response.text
        review_url = response.headers["location"]

        page = auth.get(review_url)
        assert page.status_code == 200
        assert "신규 등록" in page.text

        change_set_id = review_url.rsplit("/", 1)[-1]
        ids = re.findall(r'name="approve" value="(\d+)"', page.text)
        applied = auth.post(f"/review/{change_set_id}/apply", data={"approve": ids})
        assert applied.status_code == 303
        return auth

    def test_upload_does_not_apply_immediately(self, auth):
        upload(auth, "sample_desk_matrix.xlsx")
        assert "아직 등록된 자료가 없습니다" in auth.get("/").text

    def test_matrix_renders_after_apply(self, loaded):
        page = loaded.get("/")
        assert page.status_code == 200
        for expected in ("조선일보", "이길성", "산업부장", "편집국장"):
            assert expected in page.text

    def test_phone_is_masked_by_default(self, loaded):
        assert "●●●●" in loaded.get("/").text

    def test_category_filter(self, loaded):
        page = loaded.get("/?category=경제지")
        body = page.text.split("<tbody>", 1)[1]
        assert "매일경제" in body
        assert "조선일보" not in body

    def test_outlet_modal_fragment(self, loaded):
        page = loaded.get("/")
        outlet_id = re.search(r'data-outlet="(\d+)"', page.text).group(1)
        modal = loaded.get(f"/outlet/{outlet_id}")
        assert modal.status_code == 200
        assert "명단 전체 복사" in modal.text

    def test_search_answers_natural_question(self, loaded):
        page = loaded.get("/search", params={"q": "매경 산업부장이 누구지?"})
        assert page.status_code == 200
        assert "매일경제 · 산업부장" in page.text
        assert "정욱" in page.text

    def test_duplicate_upload_is_blocked(self, loaded):
        response = upload(loaded, "sample_desk_matrix.xlsx")
        assert response.status_code == 400
        assert "이미 등록되어" in response.text

    def test_export_returns_xlsx(self, loaded):
        response = loaded.get("/export.xlsx")
        assert response.status_code == 200
        assert response.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml"
        )
        assert response.content[:2] == b"PK"

    def test_history_lists_applied_file(self, loaded):
        page = loaded.get("/history")
        assert "sample_desk_matrix.xlsx" in page.text
        assert "반영됨" in page.text

    def test_second_file_shows_review_screen(self, loaded):
        response = upload(loaded, "sample_combined.docx")
        assert response.status_code == 303
        page = loaded.get(response.headers["location"])
        assert "변경 내역 검토" in page.text

    def test_audit_log_records_activity(self, loaded):
        page = loaded.get("/admin/users")
        assert "apply_changes" in page.text
        assert "login" in page.text


class TestLoginLockout:
    """무차별 대입 방어: 짧은 시간에 실패가 몰리면 잠근다."""

    def test_lockout_after_repeated_failures(self, client, admin):
        from app.config import get_settings

        limit = get_settings().login_fail_limit
        for _ in range(limit):
            response = client.post(
                "/login", data={"username": "tester", "password": "wrong-pass", "next": "/"}
            )
            assert response.status_code == 401
        # 한도를 넘기면 맞는 비밀번호로도 잠시 막힌다
        response = client.post(
            "/login", data={"username": "tester", "password": "verysecret123", "next": "/"}
        )
        assert response.status_code == 429
        assert "너무 많습니다" in response.text


class TestTotpLogin:
    """REQUIRE_TOTP=true 일 때의 2단계 인증 흐름 (기본값은 꺼져 있다)."""

    @pytest.fixture()
    def totp_on(self):
        from app.config import get_settings

        settings = get_settings()
        settings.require_totp = True
        yield
        settings.require_totp = False

    def test_password_alone_does_not_issue_session(self, client, admin, totp_on):
        response = client.post(
            "/login", data={"username": "tester", "password": "verysecret123", "next": "/"}
        )
        assert response.status_code == 200
        assert "2단계 인증" in response.text
        assert "contacts_session" not in response.cookies
        # 세션 없이 메인 접근 → 로그인으로 돌려보냄
        assert client.get("/").status_code == 303

    def test_enroll_and_login_with_code(self, client, admin, totp_on, db_session):
        from app.models import AppUser
        from app.services import totp as totp_service

        page = client.post(
            "/login", data={"username": "tester", "password": "verysecret123", "next": "/"}
        )
        token = re.search(r'name="token" value="([^"]+)"', page.text).group(1)
        assert "QR" in page.text  # 첫 로그인은 등록 화면

        user = db_session.get(AppUser, admin.id)
        db_session.refresh(user)
        code = totp_service.code_at(user.totp_secret)
        response = client.post("/login/otp", data={"token": token, "code": code, "next": "/"})
        assert response.status_code == 303
        assert client.get("/").status_code == 200

        db_session.refresh(user)
        assert user.totp_confirmed is True

    def test_wrong_code_is_rejected(self, client, admin, totp_on):
        page = client.post(
            "/login", data={"username": "tester", "password": "verysecret123", "next": "/"}
        )
        token = re.search(r'name="token" value="([^"]+)"', page.text).group(1)
        response = client.post("/login/otp", data={"token": token, "code": "000000", "next": "/"})
        assert response.status_code == 401
        assert "올바르지 않습니다" in response.text
        assert client.get("/").status_code == 303

    def test_totp_off_keeps_old_flow(self, client, admin):
        response = client.post(
            "/login", data={"username": "tester", "password": "verysecret123", "next": "/"}
        )
        assert response.status_code == 303  # OTP 단계 없이 바로 로그인


class TestReporterOnlyOutlets:
    """데스크 없이 출입기자만 있는 매체도 매트릭스에 행으로 나와야 한다."""

    def test_reporter_only_outlet_appears_in_matrix(self, auth):
        response = upload(auth, "sample_reporter_list.xlsx")
        assert response.status_code == 303
        review_url = response.headers["location"]
        page = auth.get(review_url)
        ids = re.findall(r'name="approve" value="(\d+)"', page.text)
        change_set_id = review_url.rsplit("/", 1)[-1]
        auth.post(f"/review/{change_set_id}/apply", data={"approve": ids})

        body = auth.get("/").text
        # 이 파일의 인원은 전부 출입기자 — 그래도 매체 행과 인원 배지가 보여야 한다
        assert "국민일보" in body
        assert "출입" in body


class TestPermissions:
    @pytest.fixture()
    def viewer(self, client, db_session):
        from app.models import AppUser
        from app.services.auth import hash_password

        db_session.add(
            AppUser(
                username="viewer1",
                display_name="조회자",
                password_hash=hash_password("verysecret123"),
                role="viewer",
            )
        )
        db_session.commit()
        client.post(
            "/login", data={"username": "viewer1", "password": "verysecret123", "next": "/"}
        )
        return client

    def test_viewer_cannot_upload(self, viewer):
        response = viewer.get("/upload")
        assert response.status_code == 403
        assert "업로드 권한이 없습니다" in response.text

    def test_viewer_cannot_manage_users(self, viewer):
        assert viewer.get("/admin/users").status_code == 403

    def test_viewer_can_view_and_search(self, viewer):
        assert viewer.get("/").status_code == 200
        assert viewer.get("/search", params={"q": "조선"}).status_code == 200


class TestUploadValidation:
    def test_rejects_wrong_extension(self, auth, tmp_path):
        bad = tmp_path / "주소록.txt"
        bad.write_text("매체,이름", encoding="utf-8")
        with bad.open("rb") as handle:
            response = auth.post("/upload", files={"file": ("주소록.txt", handle)})
        assert response.status_code == 400
        assert "xlsx 또는 docx" in response.text


class TestHighlighting:
    def test_initial_load_is_not_highlighted(self, auth):
        """처음 적재는 전부 신규이므로 변동 음영을 넣지 않는다."""
        response = upload(auth, "sample_desk_matrix.xlsx")
        review_url = response.headers["location"]
        page = auth.get(review_url)
        ids = re.findall(r'name="approve" value="(\d+)"', page.text)
        auth.post(f"{review_url}/apply", data={"approve": ids})

        body = auth.get("/").text.split("<tbody>", 1)[1]
        assert "person recent" not in body

    def test_later_change_is_highlighted(self, auth, db_session):
        """설치 당일의 연속 적재는 초기 적재, 그 뒤의 갱신부터 음영이 붙는다."""
        from datetime import timedelta

        from sqlalchemy import update

        from app.models import ChangeSet

        def load(name: str):
            response = upload(auth, name)
            review_url = response.headers["location"]
            page = auth.get(review_url)
            ids = re.findall(r'name="approve" value="(\d+)"', page.text)
            auth.post(f"{review_url}/apply", data={"approve": ids})

        load("sample_desk_matrix.xlsx")
        # 초기 적재를 하루 전으로 되돌려 '나중의 갱신' 상황을 만든다
        db_session.execute(
            update(ChangeSet).values(applied_at=ChangeSet.applied_at - timedelta(days=1))
        )
        db_session.commit()

        load("sample_combined.docx")
        body = auth.get("/").text.split("<tbody>", 1)[1]
        assert "person recent" in body


class TestEditThroughWeb:
    """화면에서 실제로 눌러 고치는 흐름."""

    @pytest.fixture()
    def loaded(self, auth):
        response = upload(auth, "sample_desk_matrix.xlsx")
        review_url = response.headers["location"]
        page = auth.get(review_url)
        ids = re.findall(r'name="approve" value="(\d+)"', page.text)
        auth.post(f"{review_url}/apply", data={"approve": ids})
        return auth

    def _first_assignment_id(self, client) -> str:
        page = client.get("/")
        outlet_id = re.search(r'data-outlet="(\d+)"', page.text).group(1)
        modal = client.get(f"/outlet/{outlet_id}")
        return re.search(r"/assignment/(\d+)/edit", modal.text).group(1)

    def test_edit_button_appears_for_editors(self, loaded):
        page = loaded.get("/")
        outlet_id = re.search(r'data-outlet="(\d+)"', page.text).group(1)
        modal = loaded.get(f"/outlet/{outlet_id}")
        assert "수정" in modal.text
        assert "+ 항목 추가" in modal.text

    def test_edit_form_renders(self, loaded):
        assignment_id = self._first_assignment_id(loaded)
        form = loaded.get(f"/assignment/{assignment_id}/edit")
        assert form.status_code == 200
        assert "항목 수정" in form.text
        assert "잘못된 항목 삭제" in form.text

    def test_saving_changes_the_matrix(self, loaded):
        assignment_id = self._first_assignment_id(loaded)
        response = loaded.post(
            f"/assignment/{assignment_id}/edit",
            data={
                "name": "고친이름",
                "phone": "010-1234-5678",
                "kind": "desk",
                "role_slot": "편집국장",
                "role_label": "편집국장",
                "dept": "",
                "note": "",
            },
        )
        assert response.status_code == 200
        assert response.headers.get("X-Contacts-Edited")
        assert "고친이름" in loaded.get("/").text

    def test_bad_phone_shows_the_form_again(self, loaded):
        assignment_id = self._first_assignment_id(loaded)
        response = loaded.post(
            f"/assignment/{assignment_id}/edit",
            data={"name": "홍길동", "phone": "12", "kind": "desk", "role_slot": "", "role_label": ""},
        )
        assert response.status_code == 400
        assert "전화번호 형식" in response.text

    def test_purge_removes_from_the_matrix(self, loaded):
        page = loaded.get("/")
        outlet_id = re.search(r'data-outlet="(\d+)"', page.text).group(1)
        modal = loaded.get(f"/outlet/{outlet_id}")
        assignment_id = re.search(r"/assignment/(\d+)/edit", modal.text).group(1)
        name = re.search(r'<a href="/person/\d+">([^<]+)</a>', modal.text).group(1)

        loaded.post(f"/assignment/{assignment_id}/delete", data={"mode": "purge"})
        after = loaded.get(f"/outlet/{outlet_id}").text
        assert f">{name}</a>" not in after

    def test_add_new_row(self, loaded):
        page = loaded.get("/")
        outlet_id = re.search(r'data-outlet="(\d+)"', page.text).group(1)
        response = loaded.post(
            f"/outlet/{outlet_id}/add",
            data={
                "name": "새로운기자",
                "phone": "010-2222-3333",
                "kind": "reporter",
                "role_slot": "",
                "role_label": "기자",
                "dept": "산업부",
                "note": "",
            },
        )
        assert response.headers.get("X-Contacts-Edited")
        assert "새로운기자" in loaded.get(f"/outlet/{outlet_id}").text

    def test_edit_is_recorded_in_the_audit_log(self, loaded):
        assignment_id = self._first_assignment_id(loaded)
        loaded.post(
            f"/assignment/{assignment_id}/edit",
            data={"name": "감사로그", "phone": "", "kind": "desk", "role_slot": "", "role_label": ""},
        )
        assert "edit_assignment" in loaded.get("/admin/users").text


class TestEditPermissions:
    @pytest.fixture()
    def viewer(self, client, db_session):
        from app.models import AppUser
        from app.services.auth import hash_password

        db_session.add(
            AppUser(
                username="viewer2",
                display_name="조회자",
                password_hash=hash_password("verysecret123"),
                role="viewer",
            )
        )
        db_session.commit()
        client.post(
            "/login", data={"username": "viewer2", "password": "verysecret123", "next": "/"}
        )
        return client

    def test_viewer_cannot_open_edit_form(self, viewer):
        assert viewer.get("/assignment/1/edit").status_code == 403

    def test_viewer_cannot_delete(self, viewer):
        assert viewer.post("/assignment/1/delete", data={"mode": "purge"}).status_code == 403

    def test_viewer_sees_no_edit_buttons(self, viewer, db_session):
        from app.models import Outlet
        from sqlalchemy import select

        outlet = db_session.scalar(select(Outlet))
        modal = viewer.get(f"/outlet/{outlet.id}")
        assert "/edit" not in modal.text


class TestEntryPage:
    """직접 입력 페이지 — 매체를 고르고 여러 명을 이어서 넣는다."""

    def _outlet_id(self, client) -> int:
        from sqlalchemy import select

        from app.models import Outlet

        page = client.get("/entry")
        assert page.status_code == 200
        return int(re.search(r'<option value="(\d+)"', page.text).group(1))

    def test_page_asks_for_an_outlet_first(self, auth):
        page = auth.get("/entry")
        assert "매체를 고르면" in page.text

    def test_nav_link_is_shown_to_editors(self, auth):
        assert "직접 입력" in auth.get("/").text

    def test_viewer_cannot_open_it(self, client, db_session):
        from app.models import AppUser
        from app.services.auth import hash_password

        db_session.add(
            AppUser(
                username="viewer3",
                display_name="조회자",
                password_hash=hash_password("verysecret123"),
                role="viewer",
            )
        )
        db_session.commit()
        client.post(
            "/login", data={"username": "viewer3", "password": "verysecret123", "next": "/"}
        )
        assert client.get("/entry").status_code == 403
        assert "직접 입력" not in client.get("/").text

    def test_add_person_with_email_and_memo(self, auth):
        outlet_id = self._outlet_id(auth)
        response = auth.post(
            "/entry",
            data={
                "outlet": str(outlet_id),
                "name": "새기자",
                "phone": "010-1234-5678",
                "email": "SaeGija@example.com",
                "kind": "reporter",
                "role_slot": "",
                "role_label": "차장",
                "dept": "산업부",
                "note": "재계",
                "memo": "행사 때 명함 교환",
            },
        )
        assert response.status_code == 303
        assert f"outlet={outlet_id}" in response.headers["location"], "같은 매체 화면에 머물러야 한다"

        page = auth.get(f"/entry?outlet={outlet_id}")
        assert "새기자" in page.text
        assert "saegija@example.com" in page.text, "이메일은 소문자로 정규화된다"
        assert "행사 때 명함 교환" in page.text

    def test_added_person_shows_in_search(self, auth):
        outlet_id = self._outlet_id(auth)
        auth.post(
            "/entry",
            data={
                "outlet": str(outlet_id),
                "name": "검색될기자",
                "phone": "010-7777-1111",
                "email": "find@example.com",
                "kind": "reporter",
                "role_slot": "",
                "role_label": "기자",
            },
        )
        assert "검색될기자" in auth.get("/search", params={"q": "검색될기자"}).text
        assert "검색될기자" in auth.get("/search", params={"q": "find@example.com"}).text

    def test_bad_email_is_rejected_with_the_form_kept(self, auth):
        outlet_id = self._outlet_id(auth)
        response = auth.post(
            "/entry",
            data={
                "outlet": str(outlet_id),
                "name": "홍길동",
                "email": "not-an-email",
                "kind": "reporter",
                "role_slot": "",
            },
        )
        assert response.status_code == 400
        assert "이메일 형식" in response.text

    def test_outlet_is_required(self, auth):
        response = auth.post("/entry", data={"name": "홍길동", "kind": "reporter"})
        assert response.status_code == 400

    def test_desk_entry_lands_in_the_matrix(self, auth):
        outlet_id = self._outlet_id(auth)
        auth.post(
            "/entry",
            data={
                "outlet": str(outlet_id),
                "name": "새논설실장",
                "phone": "010-8888-2222",
                "kind": "desk",
                "role_slot": "논설실장",
                "role_label": "논설실장",
            },
        )
        body = auth.get("/").text.split("<tbody>", 1)[1]
        assert "새논설실장" in body

    def test_email_appears_in_export(self, auth):
        outlet_id = self._outlet_id(auth)
        auth.post(
            "/entry",
            data={
                "outlet": str(outlet_id),
                "name": "내보내기기자",
                "phone": "010-3333-4444",
                "email": "export@example.com",
                "kind": "reporter",
                "role_slot": "",
            },
        )
        import io

        from openpyxl import load_workbook

        response = auth.get("/export.xlsx")
        workbook = load_workbook(io.BytesIO(response.content))
        sheet = workbook["출입기자"]
        assert sheet.cell(row=1, column=7).value == "이메일"
        values = {row[6] for row in sheet.iter_rows(min_row=2, values_only=True)}
        assert "export@example.com" in values

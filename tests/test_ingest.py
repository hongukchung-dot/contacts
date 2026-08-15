"""적재 → 비교 → 승인 → 반영 통합 테스트 (실제 PostgreSQL 사용)."""

from __future__ import annotations

import shutil
from datetime import date
from pathlib import Path

import pytest
from sqlalchemy import select

from app.models import (
    APPLIED,
    APPROVED,
    CONFLICT,
    NEW,
    PENDING,
    REMOVE,
    UPDATE,
    Assignment,
    Person,
)
from app.services.ingest import DuplicateFileError, apply_change_set, stage_file
from tests.conftest import FIXTURES, requires_db

pytestmark = requires_db


def load(session, fixture_name: str, filename: str | None = None, as_of=None):
    path = FIXTURES / fixture_name
    change_set = stage_file(
        session, path, filename=filename or fixture_name, as_of_override=as_of
    )
    session.flush()
    return change_set


def apply(session, change_set):
    counts = apply_change_set(session, change_set, applied_by_id=None)
    session.commit()
    return counts


def current_assignments(session, outlet_name: str | None = None):
    from app.models import Outlet

    stmt = select(Assignment).where(Assignment.valid_to.is_(None))
    if outlet_name:
        stmt = stmt.join(Outlet).where(Outlet.name == outlet_name)
    return session.scalars(stmt).all()


def person_by_name(session, name: str) -> Person | None:
    return session.scalar(select(Person).where(Person.name == name))


# ── 최초 적재 ───────────────────────────────────────────────────────────────

class TestFirstLoad:
    def test_all_changes_are_new_and_auto_approved(self, db_session):
        change_set = load(db_session, "sample_reporter_list.xlsx")
        assert change_set.changes
        assert {c.change_type for c in change_set.changes} == {NEW}
        assert all(c.decision == APPROVED for c in change_set.changes)

    def test_nothing_is_written_before_apply(self, db_session):
        load(db_session, "sample_reporter_list.xlsx")
        assert current_assignments(db_session) == []

    def test_apply_creates_people_and_assignments(self, db_session):
        change_set = load(db_session, "sample_reporter_list.xlsx")
        counts = apply(db_session, change_set)

        assert counts[NEW] > 0
        assert change_set.status == APPLIED
        assert person_by_name(db_session, "이길성") is not None
        chosun = current_assignments(db_session, "조선일보")
        assert {a.person.name for a in chosun} >= {"이길성", "정한국", "전수용", "안별"}

    def test_source_records_preserve_original_text(self, db_session):
        change_set = load(db_session, "sample_reporter_list.xlsx")
        records = change_set.source_file.records
        assert records
        assert all(record.source_text for record in records)
        assert change_set.source_file.as_of == date(2026, 6, 9)


# ── 재업로드 차단 ───────────────────────────────────────────────────────────

class TestDuplicate:
    def test_same_file_twice_is_rejected(self, db_session):
        change_set = load(db_session, "sample_reporter_list.xlsx")
        apply(db_session, change_set)
        with pytest.raises(DuplicateFileError):
            load(db_session, "sample_reporter_list.xlsx", filename="다시올림.xlsx")


# ── 갱신 파일 반영 ──────────────────────────────────────────────────────────

class TestUpdateCycle:
    @pytest.fixture()
    def seeded(self, db_session):
        apply(db_session, load(db_session, "sample_reporter_list.xlsx"))
        return db_session

    def test_newer_file_adds_and_updates(self, seeded):
        change_set = load(seeded, "sample_combined.docx")
        types = {c.change_type for c in change_set.changes}
        assert NEW in types
        apply(seeded, change_set)

        # docx 에만 있는 인물이 들어와야 한다
        assert person_by_name(seeded, "최인준") is not None
        assert person_by_name(seeded, "박순찬") is not None

    def test_existing_person_is_matched_by_phone_not_duplicated(self, seeded):
        apply(seeded, load(seeded, "sample_combined.docx"))
        rows = seeded.scalars(select(Person).where(Person.name == "이길성")).all()
        assert len(rows) == 1, "같은 사람이 두 번 등록되면 안 된다"

    def test_desk_records_land_in_matrix_slot(self, seeded):
        apply(seeded, load(seeded, "sample_combined.docx"))
        desks = [a for a in current_assignments(seeded, "조선일보") if a.kind == "desk"]
        assert {a.person.name for a in desks} >= {"이길성", "전수용"}
        assert all(a.role_slot == "산업부장" for a in desks)

    def test_history_is_kept_not_overwritten(self, seeded):
        apply(seeded, load(seeded, "sample_combined.docx"))
        # 출입기자로 먼저 들어온 이길성은 데스크 자리가 새로 생겨도 이력이 남는다
        person = person_by_name(seeded, "이길성")
        rows = seeded.scalars(
            select(Assignment).where(Assignment.person_id == person.id)
        ).all()
        assert len(rows) >= 1
        assert any(a.valid_to is None for a in rows)


# ── 오타 의심 탐지 ──────────────────────────────────────────────────────────

class TestTypoDetection:
    def test_one_digit_difference_becomes_conflict(self, db_session, tmp_path):
        apply(db_session, load(db_session, "sample_reporter_list.xlsx"))

        # 손철의 번호를 한 자리만 바꾼 파일을 만든다 (실제 원본에서 발견된 유형)
        modified = _rewrite_phone(
            tmp_path, "sample_reporter_list.xlsx", "010-8601-0040", "010-8601-0049"
        )
        change_set = stage_file(
            db_session, modified, filename="(26-0801) 출입기자 현황.xlsx", as_of_override=date(2026, 8, 1)
        )
        db_session.flush()

        conflicts = [c for c in change_set.changes if c.change_type == CONFLICT]
        assert conflicts, "한 자리 차이는 오타 의심으로 잡혀야 한다"
        conflict = conflicts[0]
        assert conflict.person_name == "손철"
        assert conflict.auto_apply is False
        assert conflict.decision == PENDING
        assert "한 자리" in (conflict.reason or "")

    def test_conflict_is_not_applied_without_approval(self, db_session, tmp_path):
        apply(db_session, load(db_session, "sample_reporter_list.xlsx"))
        modified = _rewrite_phone(
            tmp_path, "sample_reporter_list.xlsx", "010-8601-0040", "010-8601-0049"
        )
        change_set = stage_file(
            db_session, modified, filename="수정본.xlsx", as_of_override=date(2026, 8, 1)
        )
        db_session.flush()
        apply(db_session, change_set)

        assert person_by_name(db_session, "손철").phone == "01086010040"

    def test_conflict_is_applied_after_approval(self, db_session, tmp_path):
        apply(db_session, load(db_session, "sample_reporter_list.xlsx"))
        modified = _rewrite_phone(
            tmp_path, "sample_reporter_list.xlsx", "010-8601-0040", "010-8601-0049"
        )
        change_set = stage_file(
            db_session, modified, filename="수정본.xlsx", as_of_override=date(2026, 8, 1)
        )
        db_session.flush()
        for change in change_set.changes:
            if change.change_type == CONFLICT:
                change.decision = APPROVED
        apply(db_session, change_set)

        assert person_by_name(db_session, "손철").phone == "01086010049"

    def test_completely_different_phone_is_plain_update(self, db_session, tmp_path):
        apply(db_session, load(db_session, "sample_reporter_list.xlsx"))
        modified = _rewrite_phone(
            tmp_path, "sample_reporter_list.xlsx", "010-8601-0040", "010-2222-7777"
        )
        change_set = stage_file(
            db_session, modified, filename="수정본.xlsx", as_of_override=date(2026, 8, 1)
        )
        db_session.flush()
        phone_changes = [c for c in change_set.changes if c.field == "phone"]
        assert phone_changes
        assert phone_changes[0].change_type == UPDATE
        assert phone_changes[0].auto_apply is True


# ── 삭제 탐지 ───────────────────────────────────────────────────────────────

class TestRemoval:
    def test_missing_person_becomes_removal_needing_approval(self, db_session, tmp_path):
        apply(db_session, load(db_session, "sample_reporter_list.xlsx"))
        trimmed = _drop_row(tmp_path, "sample_reporter_list.xlsx", "정한국")

        change_set = stage_file(
            db_session, trimmed, filename="(26-0801) 출입기자 현황.xlsx", as_of_override=date(2026, 8, 1)
        )
        db_session.flush()

        removals = [c for c in change_set.changes if c.change_type == REMOVE]
        assert [c.person_name for c in removals] == ["정한국"]
        assert removals[0].auto_apply is False

    def test_removal_closes_assignment_but_keeps_history(self, db_session, tmp_path):
        apply(db_session, load(db_session, "sample_reporter_list.xlsx"))
        trimmed = _drop_row(tmp_path, "sample_reporter_list.xlsx", "정한국")
        change_set = stage_file(
            db_session, trimmed, filename="축소본.xlsx", as_of_override=date(2026, 8, 1)
        )
        db_session.flush()
        for change in change_set.changes:
            change.decision = APPROVED
        apply(db_session, change_set)

        person = person_by_name(db_session, "정한국")
        assert person is not None, "사람 자체는 지우지 않는다"
        rows = db_session.scalars(
            select(Assignment).where(Assignment.person_id == person.id)
        ).all()
        assert rows and all(a.valid_to == date(2026, 8, 1) for a in rows)

    def test_out_of_scope_outlets_are_not_removed(self, db_session, tmp_path):
        """방송 시트만 담긴 파일이 종합지 인물을 지워서는 안 된다."""
        apply(db_session, load(db_session, "sample_reporter_list.xlsx"))
        broadcast_only = _keep_sheets(tmp_path, "sample_reporter_list.xlsx", {"방송"})

        change_set = stage_file(
            db_session, broadcast_only, filename="방송만.xlsx", as_of_override=date(2026, 8, 1)
        )
        db_session.flush()
        removals = [c for c in change_set.changes if c.change_type == REMOVE]
        assert not [c for c in removals if c.outlet_name == "조선일보"]


# ── 오래된 파일 방어 ────────────────────────────────────────────────────────

class TestStaleFile:
    def test_older_file_disables_auto_apply(self, db_session, tmp_path):
        apply(db_session, load(db_session, "sample_combined.docx"))  # 8/3

        older = tmp_path / "옛날파일.xlsx"
        shutil.copy(FIXTURES / "sample_reporter_list.xlsx", older)
        change_set = stage_file(
            db_session, older, filename="(26-0609) 출입기자 현황.xlsx", as_of_override=date(2026, 6, 9)
        )
        db_session.flush()

        assert change_set.changes
        assert all(not c.auto_apply for c in change_set.changes)
        assert all(c.decision == PENDING for c in change_set.changes)


# ── 헬퍼: 픽스처 변형 ───────────────────────────────────────────────────────

def _rewrite_phone(tmp_path: Path, fixture: str, old: str, new: str) -> Path:
    from openpyxl import load_workbook

    target = tmp_path / f"mod_{fixture}"
    shutil.copy(FIXTURES / fixture, target)
    workbook = load_workbook(target)
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and old in cell.value:
                    cell.value = cell.value.replace(old, new)
    workbook.save(target)
    return target


def _drop_row(tmp_path: Path, fixture: str, name: str) -> Path:
    from openpyxl import load_workbook

    target = tmp_path / f"drop_{fixture}"
    shutil.copy(FIXTURES / fixture, target)
    workbook = load_workbook(target)
    for sheet in workbook.worksheets:
        for idx in range(sheet.max_row, 0, -1):
            values = [c.value for c in sheet[idx]]
            if any(isinstance(v, str) and v.strip() == name for v in values):
                sheet.delete_rows(idx)
    workbook.save(target)
    return target


def _keep_sheets(tmp_path: Path, fixture: str, keep: set[str]) -> Path:
    from openpyxl import load_workbook

    target = tmp_path / f"keep_{fixture}"
    shutil.copy(FIXTURES / fixture, target)
    workbook = load_workbook(target)
    for title in [s.title for s in workbook.worksheets if s.title not in keep]:
        del workbook[title]
    workbook.save(target)
    return target


class TestRemovalFalsePositives:
    """부서·구분이 바뀐 것뿐인 사람을 삭제로 잡으면 안 된다."""

    @pytest.fixture()
    def seeded(self, db_session):
        apply(db_session, load(db_session, "sample_reporter_list.xlsx"))
        return db_session

    def test_reporter_promoted_to_desk_is_not_a_removal(self, seeded):
        """이길성은 6/9 파일엔 출입기자, 8/3 파일엔 데스크로 나온다."""
        change_set = load(seeded, "sample_combined.docx")
        removed = {c.person_name for c in change_set.changes if c.change_type == REMOVE}
        assert "이길성" not in removed
        assert "전수용" not in removed
        assert "정한국" not in removed

    def test_dept_change_is_not_a_removal(self, seeded):
        change_set = load(seeded, "sample_combined.docx")
        removed = {c.person_name for c in change_set.changes if c.change_type == REMOVE}
        assert "석민수" not in removed
        assert "박예원" not in removed

    def test_genuinely_absent_person_is_still_flagged(self, seeded):
        """반대로 정말 빠진 사람은 계속 잡아야 한다."""
        apply(seeded, load(seeded, "sample_desk_matrix.xlsx"))
        change_set = load(seeded, "sample_combined.docx")
        removed = {c.person_name for c in change_set.changes if c.change_type == REMOVE}
        assert "신수정" in removed  # 동아일보 산업2부장 — docx 에 없음
